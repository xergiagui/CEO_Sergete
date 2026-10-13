"""
Registro Excel de un proyecto: UNA FILA por clip y plataforma.

El Excel es la fuente de verdad: el script lo lee, añade/actualiza filas y lo
guarda, respetando lo que Sergio haya tocado a mano (Estado, Fecha
publicación, Notas, columnas propias). Las columnas se localizan por su
CABECERA, así que reordenarlas o añadir columnas propias no rompe nada.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import config

HOJA = "Registro"
COLS_ANTES = ["ID", "Clip", "Plataforma", "Estado", "Tema"]
COLS_DESPUES = ["Título", "Copy", "Hashtags", "Texto para pegar", "Nº caracteres",
                "Fecha alta", "Fecha publicación", "Notas", "Ruta"]
ANCHOS = {"ID": 10, "Clip": 34, "Plataforma": 14, "Estado": 14, "Tema": 32, "Título": 32,
          "Copy": 60, "Hashtags": 42, "Texto para pegar": 60, "Nº caracteres": 10,
          "Fecha alta": 12, "Fecha publicación": 13, "Notas": 30, "Ruta": 60}
COLS_TEXTO_LARGO = {"Tema", "Título", "Copy", "Hashtags", "Texto para pegar", "Notas"}
COLS_FECHA = {"Fecha alta", "Fecha publicación"}


class RegistroBloqueado(Exception):
    pass


def columnas(proyecto) -> list[str]:
    return COLS_ANTES + [c["columna"] for c in proyecto.campos_extra()] + COLS_DESPUES


def esta_bloqueado(path: Path) -> bool:
    """True si el Excel está abierto (Excel bloquea la escritura en Windows)."""
    if not path.exists():
        return False
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True


class Registro:
    def __init__(self, path: Path, proyecto):
        self.path = Path(path)
        self.proyecto = proyecto
        if self.path.exists():
            try:
                self.wb = load_workbook(self.path)
            except PermissionError:
                raise RegistroBloqueado(
                    f"El Excel está abierto o bloqueado: {self.path}\nCiérralo y repite la orden (no se ha escrito nada).")
            self.ws = self.wb[HOJA] if HOJA in self.wb.sheetnames else self.wb.active
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = HOJA
            self._crear_cabecera()
        self._mapear_columnas()

    # ---------- estructura ----------
    def _crear_cabecera(self) -> None:
        ws = self.ws
        cols = columnas(self.proyecto)
        relleno = PatternFill("solid", fgColor="1F2937")
        for i, nombre in enumerate(cols, start=1):
            c = ws.cell(row=1, column=i, value=nombre)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = relleno
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[c.column_letter].width = ANCHOS.get(nombre, 22)
        ws.freeze_panes = "C2"
        ws.row_dimensions[1].height = 22
        estado_col = ws.cell(row=1, column=cols.index("Estado") + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(config.ESTADOS) + '"', allow_blank=True)
        dv.add(f"{estado_col}2:{estado_col}5000")
        ws.add_data_validation(dv)

    def _mapear_columnas(self) -> None:
        self.col = {}
        for c in self.ws[1]:
            if c.value:
                self.col[str(c.value).strip()] = c.column
        # Columnas que falten (p. ej. un campo propio añadido después) → al final.
        for nombre in columnas(self.proyecto):
            if nombre not in self.col:
                idx = self.ws.max_column + 1
                c = self.ws.cell(row=1, column=idx, value=nombre)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F2937")
                self.ws.column_dimensions[c.column_letter].width = ANCHOS.get(nombre, 22)
                self.col[nombre] = idx

    # ---------- lectura ----------
    def filas(self):
        """[(nº_fila, {columna: valor})] de las filas con datos."""
        out = []
        for r in range(2, self.ws.max_row + 1):
            d = {n: self.ws.cell(row=r, column=i).value for n, i in self.col.items()}
            if d.get("ID") or d.get("Clip"):
                out.append((r, d))
        return out

    def buscar(self, clip: str, plataforma: str) -> int | None:
        clip_l, plat_l = clip.lower(), plataforma.lower()
        for r, d in self.filas():
            if (str(d.get("Clip") or "").lower() == clip_l
                    and str(d.get("Plataforma") or "").lower() == plat_l):
                return r
        return None

    def id_de_clip(self, clip: str) -> str | None:
        for _, d in self.filas():
            if str(d.get("Clip") or "").lower() == clip.lower() and d.get("ID"):
                return d["ID"]
        return None

    def siguiente_id(self) -> str:
        pref = self.proyecto.prefijo_id + "-"
        nums = []
        for _, d in self.filas():
            v = str(d.get("ID") or "")
            if v.startswith(pref) and v[len(pref):].isdigit():
                nums.append(int(v[len(pref):]))
        return f"{pref}{(max(nums) if nums else 0) + 1:04d}"

    # ---------- escritura ----------
    def escribir(self, fila: int | None, datos: dict) -> int:
        """Escribe `datos` (columna → valor) en `fila` o en una nueva. Devuelve la fila."""
        if fila is None:
            fila = max([r for r, _ in self.filas()] or [1]) + 1
        for nombre, valor in datos.items():
            if nombre not in self.col:
                continue
            c = self.ws.cell(row=fila, column=self.col[nombre], value=valor)
            if nombre in COLS_TEXTO_LARGO:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                c.alignment = Alignment(vertical="top")
            if nombre in COLS_FECHA and isinstance(valor, (dt.date, dt.datetime)):
                c.number_format = "dd/mm/yyyy"
        return fila

    def enlazar(self, fila: int, ruta: Path | None) -> None:
        """Clip con hipervínculo al MP4 (clic = se abre) + ruta en texto."""
        c = self.ws.cell(row=fila, column=self.col["Clip"])
        r = self.ws.cell(row=fila, column=self.col["Ruta"])
        if ruta and Path(ruta).exists():
            c.hyperlink = Path(ruta).resolve().as_uri()
            c.font = Font(color="0563C1", underline="single")
            r.value = str(Path(ruta).resolve())
        else:
            c.hyperlink = None
            c.font = Font()
            if ruta:
                r.value = str(ruta)
        c.alignment = Alignment(vertical="top")
        r.alignment = Alignment(vertical="top")

    def guardar(self) -> None:
        if esta_bloqueado(self.path):
            raise RegistroBloqueado(
                f"El Excel está abierto: {self.path}\nCiérralo y vuelve a lanzar la orden (no se ha escrito nada).")
        self.ws.auto_filter.ref = f"A1:{self.ws.cell(row=1, column=self.ws.max_column).column_letter}{max(self.ws.max_row, 2)}"
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(self.path)
