"""
Registro Excel de un proyecto: UNA FILA por pieza (Reel o carrusel) y plataforma.

El Excel es la fuente de verdad: el script lo lee, añade/actualiza filas y lo
guarda, respetando lo que Sergio haya tocado a mano (Estado, Fecha prevista,
Hora prevista, Fecha publicación, Notas, columnas propias). Las columnas se
localizan por su CABECERA, así que reordenarlas o añadir columnas propias no
rompe nada; a un Excel antiguo se le AÑADEN al final las columnas que le falten,
sin tocar lo escrito.

Hojas:
  Registro      una fila por pieza y plataforma (la que se edita a mano).
  Historias     plan de historias de IG por día (lo rehace `calendario planificar`;
                las fechas que cambies a mano se respetan).
  _calendario   (oculta) fechas que puso el planificador. Sirve para distinguir las
                fechas MANUALES (las que no coinciden con esta hoja), que nunca se mueven.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import config
from lib.util import normalizar

HOJA = "Registro"
HOJA_HISTORIAS = "Historias"
HOJA_PLAN = "_calendario"
COLS_ANTES = ["ID", "Clip", "Formato", "Plataforma", "Estado", "Fecha prevista", "Hora prevista", "Tema"]
COLS_DESPUES = ["Título", "Copy", "Hashtags", "Texto para pegar", "Nº caracteres",
                "Fecha alta", "Fecha publicación", "Notas", "Ruta"]
ANCHOS = {"ID": 10, "Clip": 34, "Formato": 11, "Plataforma": 14, "Estado": 14, "Fecha prevista": 13,
          "Hora prevista": 9, "Tema": 32, "Objetivo": 13, "Pilar": 13, "Título": 32,
          "Copy": 60, "Hashtags": 42, "Texto para pegar": 60, "Nº caracteres": 10,
          "Fecha alta": 12, "Fecha publicación": 13, "Notas": 30, "Ruta": 60}
COLS_TEXTO_LARGO = {"Tema", "Título", "Copy", "Hashtags", "Texto para pegar", "Notas"}
COLS_FECHA = {"Fecha alta", "Fecha publicación", "Fecha prevista"}
COLS_HORA = {"Hora prevista"}
COLS_HISTORIAS = ["Fecha prevista", "Archivo", "Carpeta"]
COLS_PLAN = ["ID", "Fecha", "Hora", "Asignado"]
_CABECERA_FONDO = "1F2937"


class RegistroBloqueado(Exception):
    pass


def columnas(proyecto) -> list[str]:
    return COLS_ANTES + [c["columna"] for c in proyecto.campos_extra()] + COLS_DESPUES


def listas(proyecto) -> dict[str, list[str]]:
    """Columnas con desplegable: Estado, Formato y los campos propios con lista de valores."""
    out = {"Estado": list(config.ESTADOS), "Formato": list(config.FORMATOS.values())}
    for campo in proyecto.campos_extra():
        if isinstance(campo.get("valores"), list):
            out[campo["columna"]] = [str(v) for v in campo["valores"]]
    return out


def formato_de_fila(d: dict) -> str:
    """'reel' | 'carrusel' de una fila (columna Formato; si está vacía, se deduce del Clip)."""
    v = normalizar(str(d.get("Formato") or ""))
    for k, nombre in config.FORMATOS.items():
        if v and v in (k, normalizar(nombre)):
            return k
    return "reel" if Path(str(d.get("Clip") or "")).suffix.lower() in config.EXT_VIDEO else "carrusel"


def esta_bloqueado(path: Path) -> bool:
    """True si el Excel está abierto (Excel bloquea la escritura en Windows)."""
    if not path.exists():
        return False
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True


def _estilo_cabecera(c) -> None:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=_CABECERA_FONDO)
    c.alignment = Alignment(vertical="center", wrap_text=True)


class Registro:
    def __init__(self, path: Path, proyecto):
        self.path = Path(path)
        self.proyecto = proyecto
        if self.path.exists():
            try:
                self.wb = load_workbook(self.path)
            except PermissionError:
                raise RegistroBloqueado(
                    f"El Excel está abierto o bloqueado: {self.path}\nCiérralo y repite la orden (no se ha escrito nada).")
            self.ws = self.wb[HOJA] if HOJA in self.wb.sheetnames else self.wb.active
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = HOJA
            self._crear_cabecera()
        self._mapear_columnas()

    # ---------- estructura ----------
    def _crear_cabecera(self) -> None:
        ws = self.ws
        for i, nombre in enumerate(columnas(self.proyecto), start=1):
            c = ws.cell(row=1, column=i, value=nombre)
            _estilo_cabecera(c)
            ws.column_dimensions[c.column_letter].width = ANCHOS.get(nombre, 22)
        ws.freeze_panes = "C2"
        ws.row_dimensions[1].height = 22
        for nombre, valores in listas(self.proyecto).items():
            self._desplegable(columnas(self.proyecto).index(nombre) + 1, valores)

    def _desplegable(self, col: int, valores: list[str]) -> None:
        formula = '"' + ",".join(valores) + '"'
        if "," in "".join(valores) or len(formula) > 255:
            return  # Excel no admite comas dentro de los valores ni listas tan largas
        letra = self.ws.cell(row=1, column=col).column_letter
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(f"{letra}2:{letra}5000")
        self.ws.add_data_validation(dv)

    def _mapear_columnas(self) -> None:
        self.col = {}
        for c in self.ws[1]:
            if c.value:
                self.col[str(c.value).strip()] = c.column
        # Columnas que falten (Excel antiguo o campo propio añadido después) → al final,
        # sin tocar las que ya hay ni lo escrito en ellas.
        desplegables = listas(self.proyecto)
        self.columnas_nuevas: list[str] = []
        for nombre in columnas(self.proyecto):
            if nombre not in self.col:
                idx = self.ws.max_column + 1
                c = self.ws.cell(row=1, column=idx, value=nombre)
                _estilo_cabecera(c)
                self.ws.column_dimensions[c.column_letter].width = ANCHOS.get(nombre, 22)
                self.col[nombre] = idx
                self.columnas_nuevas.append(nombre)
                if nombre in desplegables:
                    self._desplegable(idx, desplegables[nombre])

    # ---------- lectura ----------
    def filas(self):
        """[(nº_fila, {columna: valor})] de las filas con datos."""
        out = []
        for r in range(2, self.ws.max_row + 1):
            d = {n: self.ws.cell(row=r, column=i).value for n, i in self.col.items()}
            if d.get("ID") or d.get("Clip"):
                out.append((r, d))
        return out

    def buscar(self, clip: str, plataforma: str) -> int | None:
        clip_l, plat_l = clip.lower(), plataforma.lower()
        for r, d in self.filas():
            if (str(d.get("Clip") or "").lower() == clip_l
                    and str(d.get("Plataforma") or "").lower() == plat_l):
                return r
        return None

    def id_de_clip(self, clip: str) -> str | None:
        for _, d in self.filas():
            if str(d.get("Clip") or "").lower() == clip.lower() and d.get("ID"):
                return d["ID"]
        return None

    def siguiente_id(self) -> str:
        """Mismo contador para Reels y carruseles (MM-0001, MM-0002…)."""
        pref = self.proyecto.prefijo_id + "-"
        nums = []
        for _, d in self.filas():
            v = str(d.get("ID") or "")
            if v.startswith(pref) and v[len(pref):].isdigit():
                nums.append(int(v[len(pref):]))
        return f"{pref}{(max(nums) if nums else 0) + 1:04d}"

    # ---------- escritura ----------
    def escribir(self, fila: int | None, datos: dict) -> int:
        """Escribe `datos` (columna → valor) en `fila` o en una nueva. Devuelve la fila."""
        if fila is None:
            fila = max([r for r, _ in self.filas()] or [1]) + 1
        for nombre, valor in datos.items():
            if nombre not in self.col:
                continue
            c = self.ws.cell(row=fila, column=self.col[nombre], value=valor)
            if nombre in COLS_TEXTO_LARGO:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                c.alignment = Alignment(vertical="top")
            if nombre in COLS_FECHA and isinstance(valor, (dt.date, dt.datetime)):
                c.number_format = "dd/mm/yyyy"
            if nombre in COLS_HORA and isinstance(valor, dt.time):
                c.number_format = "hh:mm"
        return fila

    def enlazar(self, fila: int, ruta: Path | None) -> None:
        """Clip con hipervínculo al MP4 o a la carpeta del carrusel (clic = se abre) + ruta en texto."""
        c = self.ws.cell(row=fila, column=self.col["Clip"])
        r = self.ws.cell(row=fila, column=self.col["Ruta"])
        if ruta and Path(ruta).exists():
            c.hyperlink = Path(ruta).resolve().as_uri()
            c.font = Font(color="0563C1", underline="single")
            r.value = str(Path(ruta).resolve())
        else:
            c.hyperlink = None
            c.font = Font()
            if ruta:
                r.value = str(ruta)
        c.alignment = Alignment(vertical="top")
        r.alignment = Alignment(vertical="top")

    # ---------- hojas del calendario ----------
    def _hoja(self, nombre: str, cabecera: list[str], oculta: bool = False):
        if nombre in self.wb.sheetnames:
            return self.wb[nombre]
        ws = self.wb.create_sheet(nombre)
        for i, n in enumerate(cabecera, start=1):
            c = ws.cell(row=1, column=i, value=n)
            _estilo_cabecera(c)
            ws.column_dimensions[c.column_letter].width = {"Archivo": 48, "Carpeta": 60, "Asignado": 20}.get(n, 14)
        ws.freeze_panes = "A2"
        if oculta:
            ws.sheet_state = "hidden"
        return ws

    @staticmethod
    def _leer_hoja(ws) -> list[dict]:
        cab = {str(c.value).strip(): c.column for c in ws[1] if c.value}
        out = []
        for r in range(2, ws.max_row + 1):
            d = {n: ws.cell(row=r, column=i).value for n, i in cab.items()}
            if any(v not in (None, "") for v in d.values()):
                out.append(d)
        return out

    def historias(self) -> list[dict]:
        """[{'Fecha prevista', 'Archivo', 'Carpeta'}] tal cual (fechas sin interpretar)."""
        if HOJA_HISTORIAS not in self.wb.sheetnames:
            return []
        return [d for d in self._leer_hoja(self.wb[HOJA_HISTORIAS]) if d.get("Archivo")]

    def escribir_historias(self, filas: list[tuple[dt.date, str, str]]) -> None:
        """Rehace la hoja Historias: [(fecha, archivo, carpeta)]."""
        ws = self._hoja(HOJA_HISTORIAS, COLS_HISTORIAS)
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for i, (fecha, archivo, carpeta) in enumerate(sorted(filas, key=lambda f: (f[0], f[1].lower())), start=2):
            c = ws.cell(row=i, column=1, value=fecha)
            c.number_format = "dd/mm/yyyy"
            ws.cell(row=i, column=2, value=archivo)
            ws.cell(row=i, column=3, value=carpeta)

    def plan_auto(self) -> dict[str, tuple[str, str]]:
        """{ID: ('AAAA-MM-DD', 'HH:MM')} de las fechas que puso el planificador."""
        if HOJA_PLAN not in self.wb.sheetnames:
            return {}
        return {str(d["ID"]): (str(d.get("Fecha") or ""), str(d.get("Hora") or ""))
                for d in self._leer_hoja(self.wb[HOJA_PLAN]) if d.get("ID")}

    def escribir_plan_auto(self, plan: dict[str, tuple[str, str]]) -> None:
        ws = self._hoja(HOJA_PLAN, COLS_PLAN, oculta=True)
        antes = {str(d["ID"]): d for d in self._leer_hoja(ws) if d.get("ID")}
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        ahora = dt.datetime.now().isoformat(timespec="seconds")
        for i, (pid, (fecha, hora)) in enumerate(sorted(plan.items()), start=2):
            previo = antes.get(pid) or {}
            igual = (str(previo.get("Fecha") or ""), str(previo.get("Hora") or "")) == (fecha, hora)
            asignado = previo.get("Asignado") if igual and previo.get("Asignado") else ahora
            for j, v in enumerate((pid, fecha, hora, asignado), start=1):
                ws.cell(row=i, column=j, value=v)

    def guardar(self) -> None:
        if esta_bloqueado(self.path):
            raise RegistroBloqueado(
                f"El Excel está abierto: {self.path}\nCiérralo y vuelve a lanzar la orden (no se ha escrito nada).")
        self.ws.auto_filter.ref = f"A1:{self.ws.cell(row=1, column=self.ws.max_column).column_letter}{max(self.ws.max_row, 2)}"
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(self.path)
